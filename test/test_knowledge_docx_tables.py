"""DOCX ingestion must retain table text in its surrounding document order."""

from pathlib import Path

import pytest

from kiro_crew.knowledge.readers import FileReader

docx = pytest.importorskip("docx")


def test_table_only_document_retains_cell_text(tmp_path: Path) -> None:
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table._cells, ["Product", "Price", "Widget", "$42"]):
        cell.text = text
    path = tmp_path / "price-list.docx"
    doc.save(path)

    text, meta = FileReader().read(str(path))

    assert text == "Product\nPrice\nWidget\n$42"
    assert meta["format"] == "docx"
    assert meta["content_type"] == "markdown"
    assert meta["line_count"] == 4


def test_table_paragraphs_keep_body_order_and_nested_content(tmp_path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Inventory", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First cell"
    table.cell(0, 0).add_paragraph("Second paragraph")
    cell = table.cell(0, 1)
    cell.text = "Before nested table"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested value"
    cell.add_paragraph("After nested table")
    doc.add_paragraph("After outer table")
    path = tmp_path / "inventory.docx"
    doc.save(path)

    text, meta = FileReader().read(str(path))

    assert text.splitlines() == [
        "# Inventory",
        "First cell",
        "Second paragraph",
        "Before nested table",
        "Nested value",
        "",
        "After nested table",
        "After outer table",
    ]
    assert meta["paragraph_count"] == 2


def test_merged_cell_text_is_not_repeated(tmp_path: Path) -> None:
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "Shared decision"
    path = tmp_path / "merged.docx"
    doc.save(path)

    text, _ = FileReader().read(str(path))

    # Vertical-merge continuation cells may carry empty physical paragraphs.
    assert text.strip() == "Shared decision"


def test_content_control_wrapped_paragraphs_are_retained(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement

    doc = docx.Document()
    doc.add_paragraph("Plain paragraph")
    # Wrap a paragraph in a block content control: w:sdt > w:sdtContent > w:p,
    # exactly as Word emits from a template field.
    para = doc.add_paragraph("Content-control paragraph")
    sdt = OxmlElement("w:sdt")
    sdt_content = OxmlElement("w:sdtContent")
    para._p.addprevious(sdt)
    sdt.append(sdt_content)
    sdt_content.append(para._p)
    path = tmp_path / "content-control.docx"
    doc.save(path)

    text, _ = FileReader().read(str(path))

    assert text == "Plain paragraph\nContent-control paragraph"


def test_revision_wrapped_paragraphs_remain_excluded(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement

    doc = docx.Document()
    doc.add_paragraph("Visible paragraph")
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "Visible cell"
    for parent in (doc, cell):
        para = parent.add_paragraph("Unaccepted insertion")
        wrapper = OxmlElement("w:ins")
        para._p.addprevious(wrapper)
        wrapper.append(para._p)
    path = tmp_path / "revisions.docx"
    doc.save(path)

    text, _ = FileReader().read(str(path))

    assert text == "Visible paragraph\nVisible cell"
